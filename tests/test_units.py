"""단위 테스트 — DB·서버 없이 도는 순수 함수 검증."""
from __future__ import annotations

import pytest

from app.vector_search import _ctx_stem, _vec_to_str, to_contains_query


class TestToContainsQuery:
    """자연어 → Oracle Text ACCUM 구문 변환.

    회귀 가드: 2026-09-04 이전에는 사용자 문장을 CONTAINS 에 그대로 넣어
    ORA-29902 로 터졌고, LIKE 폴백이 0건이라 하이브리드가 벡터 전용으로 퇴화했다.
    """

    def test_자연어_문장이_ACCUM_구문으로_변환된다(self):
        got = to_contains_query("인덱스를 효율적으로 사용하려면 어떻게 SQL을 작성해야 하나요?")
        assert got is not None
        assert "?" not in got, "예약 연산자가 남으면 ORA-29902 가 난다"
        assert "," in got, "AND 가 아니라 ACCUM(쉼표)이어야 한다"
        assert "인덱스%" in got, "조사를 떼고 우측 절단해야 한다"

    def test_영문에는_우측절단을_붙이지_않는다(self):
        # "SQL을" 을 우측 절단하면 "SQ%" 로 과매칭된다
        got = to_contains_query("SQL을 작성한다")
        assert "SQL" in got
        assert "SQ%" not in got

    def test_질문_어투는_불용어로_빠진다(self):
        got = to_contains_query("인덱스가 무엇인가요")
        assert "무엇" not in got

    @pytest.mark.parametrize("bad", ["", "?!?", "   ", "&|~"])
    def test_쓸모없는_입력은_None(self, bad):
        # None 이면 호출부가 LIKE 로 폴백한다
        assert to_contains_query(bad) is None

    def test_중복_어간은_한_번만(self):
        got = to_contains_query("인덱스 인덱스를 인덱스가")
        assert got.count("인덱스") == 1


class TestCtxStem:
    @pytest.mark.parametrize("token,expected", [
        ("인덱스를", "인덱스"),
        ("효율적으로", "효율"),
        ("SQL을", "SQL"),
        ("인덱스", "인덱스"),
        ("SELECT", "SELECT"),
    ])
    def test_어간_추출(self, token, expected):
        assert _ctx_stem(token) == expected


class TestVecToStr:
    def test_벡터를_오라클_리터럴로(self):
        assert _vec_to_str([1.0, -0.5, 0.25]) == "[1.0,-0.5,0.25]"

    def test_공백이_없어야_한다(self):
        # 공백이 섞이면 TO_VECTOR 파싱이 실패할 수 있다
        assert " " not in _vec_to_str([0.1, 0.2, 0.3])


class TestSlidesCatalog:
    """장표 연동(P5, 2026-09-29) — 꼬리표 규칙과 앵커 풀기. 덱 파일 없이 순수 함수만."""

    def test_꼬리표_형식(self):
        from app.slides import is_tag
        assert is_tag("VS-12") and is_tag("SA-003") and is_tag("OVW-01")
        assert not is_tag("vs-12") and not is_tag("VS12") and not is_tag("V-1") and not is_tag("VSXX-12")

    def test_레지스트리_slides_는_전부_꼬리표_형식이다(self):
        from app.feature_registry import FEATURES
        from app.slides import is_tag
        for f in FEATURES:
            assert isinstance(f["slides"], list)
            for t in f["slides"]:
                assert is_tag(t), f"{f['name']} 의 slides 에 꼬리표 형식이 아닌 값: {t!r}"

    def test_앵커는_덱에_있는_꼬리표만_풀리고_나머지는_unresolved(self, monkeypatch):
        from app import slides as mod
        fake_features = [
            {"tab": "vector", "tab_label": "AI Vector Search", "name": "A", "path": "/vector?sub=search", "slides": ["VS-01", "VS-99"]},
            {"tab": "graph", "tab_label": "Property Graph", "name": "B", "path": "/graph", "slides": []},
        ]
        monkeypatch.setattr(mod, "FEATURES", fake_features)
        decks = [{"deck": "vector", "title": "벡터", "tags": {"VS-01": 3}}]
        anchors, unresolved = mod.resolve_anchors(decks)
        assert [(a["tag"], a["deck"], a["page"], a["feature"]) for a in anchors] == [("VS-01", "vector", 3, "A")]
        assert [u["tag"] for u in unresolved] == ["VS-99"]


class TestSelectAiLogSummary:
    """run_select_ai 가 이력에 남길 요약 — 액션별로 무엇을 generated_sql / response_text / row_count 로 보내는가 (PoC 1-A)."""

    def test_showsql_은_generated_sql(self):
        from app.select_ai import _summarize_result
        assert _summarize_result("showsql", "SELECT 1 FROM dual") == ("SELECT 1 FROM dual", None, None)

    def test_runsql_은_행수와_JSON_앞부분(self):
        from app.select_ai import _summarize_result
        gen, resp, rows = _summarize_result("runsql", '[{"A": 1}, {"A": 2}]')
        assert gen is None and rows == 2 and resp.startswith("[")

    def test_runsql_이_JSON_이_아니면_텍스트(self):
        from app.select_ai import _summarize_result
        assert _summarize_result("runsql", "no rows")[1] == "no rows"

    def test_없는_결과는_전부_None(self):
        from app.select_ai import _summarize_result
        assert _summarize_result("chat", None) == (None, None, None)


class TestFewshotParse:
    """PoC 2-A — CSV/JSON/XLSX 파서. 헤더 이름은 느슨하게(question/질문, sql/SQL, note/설명)."""

    def test_csv_한글_헤더와_BOM(self):
        from app.fewshot import parse_file
        content = "\ufeff질문,SQL,설명\n채널 수는?,SELECT COUNT(*) FROM channels,채널 표\n,SELECT 1 FROM dual,\n".encode("utf-8")
        rows, headers = parse_file("a.csv", content)
        assert headers == ["질문", "SQL", "설명"]
        assert rows[0]["question"] == "채널 수는?" and rows[0]["sql"] == "SELECT COUNT(*) FROM channels" and rows[0]["note"] == "채널 표" and rows[0]["error"] is None
        assert rows[1]["error"] and "질문" in rows[1]["error"]

    def test_json_배열과_rows_객체(self):
        from app.fewshot import parse_file
        rows, _ = parse_file("a.json", b'[{"question": "q1", "sql": "SELECT 1 FROM dual;"}]')
        assert rows[0]["sql"] == "SELECT 1 FROM dual"      # 끝 세미콜론 제거
        rows, _ = parse_file("b.json", b'{"rows": [{"prompt": "q2", "query": "SELECT 2 FROM dual", "comment": "n"}]}')
        assert rows[0]["question"] == "q2" and rows[0]["note"] == "n"

    def test_xlsx(self):
        import io

        import openpyxl

        from app.fewshot import parse_file
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["question", "sql", "note"])
        ws.append(["q", "SELECT 3 FROM dual", None])
        ws.append([None, None, None])
        buf = io.BytesIO()
        wb.save(buf)
        rows, headers = parse_file("a.xlsx", buf.getvalue())
        assert headers == ["question", "sql", "note"] and len(rows) == 1 and rows[0]["sql"] == "SELECT 3 FROM dual"


class TestProfileWizard:
    """PoC 3-A — 폼 → CREATE_PROFILE attributes. OCI 는 region 필수, 빈 값은 안 넣는다."""

    def test_openai_호환(self):
        from app.profiles import build_attributes, plsql_for
        attrs, errs = build_attributes({"provider": "openai", "credential_name": "GEMINI_CRED", "model": "gemini-3.8-flash",
                                        "provider_endpoint": "https://generativelanguage.googleapis.com/v1beta/openai",
                                        "object_list": [{"owner": "admin", "name": "poc_stores"}], "annotations": True, "conversation": False, "embedding_model": ""})
        assert not errs and attrs["object_list"] == [{"owner": "ADMIN", "name": "POC_STORES"}] and attrs["annotations"] is True and "embedding_model" not in attrs
        assert "DBMS_CLOUD_AI.CREATE_PROFILE" in plsql_for("X", attrs)

    def test_oci_는_region_필수(self):
        from app.profiles import build_attributes, oci_host
        _, errs = build_attributes({"provider": "oci", "credential_name": "OCI_CRED", "model": "cohere.command-r-plus-08-2024"})
        assert any("region" in e for e in errs)
        attrs, errs = build_attributes({"provider": "oci", "credential_name": "OCI_CRED", "model": "cohere.command-r-plus-08-2024", "region": "ap-osaka-1", "oci_apiformat": "COHERE"})
        assert not errs and attrs["region"] == "ap-osaka-1" and attrs["oci_apiformat"] == "COHERE"
        assert oci_host("ap-osaka-1") == "inference.generativeai.ap-osaka-1.oci.oraclecloud.com"

    def test_크리덴셜_유형_추정(self):
        from app.profiles import credential_kind, validate_name
        assert credential_kind("OCI$RESOURCE_PRINCIPAL", None) == "resource_principal"
        assert credential_kind("OCI_KEY", "ocid1.user.oc1..aaaa") == "oci_api_key"
        assert credential_kind("GEMINI_CRED", "GEMINI") == "api_key"
        assert validate_name("MY_PROFILE") is None and validate_name("1bad") and validate_name("has space")

